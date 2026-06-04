import io, math
import numpy as np
import pandas as pd
import streamlit as st
import matplotlib.pyplot as plt
import holidays
from ortools.sat.python import cp_model
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference

st.set_page_config(page_title="WFM · Pronóstico y Roster", layout="wide")
st.title("📞 WFM — Pronóstico, Dimensionamiento y Roster")

NOM = ["Lun", "Mar", "Mié", "Jue", "Vie", "Sáb", "Dom"]
LIBRES = {p: {p, (p + 1) % 7} for p in range(7)}


def erlang_b(a, c):
    b = 1.0
    for k in range(1, a + 1):
        b = (c * b) / (k + c * b)
    return b


def erlang_c(a, c):
    b = erlang_b(a, c); r = c / a
    return b / (1 - r + r * b)


def nivel_servicio(a, c, aht, asa):
    return 0.0 if a <= c else 1 - erlang_c(a, c) * math.exp(-(a - c) * asa / aht)


def nivel_atencion(ag, carga, aht, pac, NMAX=250):
    # Erlang A (con abandono): fracción de llamadas atendidas (no abandonadas)
    if carga <= 0:
        return 1.0
    lam = carga / aht; mu = 1.0 / aht; theta = 1.0 / pac
    p = [1.0]
    for n in range(1, NMAX + 1):
        muerte = n * mu if n <= ag else ag * mu + (n - ag) * theta
        p.append(p[-1] * lam / muerte)
    S = sum(p)
    aband = sum(p[n] * max(0, (n - ag)) * theta for n in range(len(p))) / S
    return 1 - aband / lam


# ---------------- Selector de vista ----------------
vista = st.sidebar.radio("Vista", ["Planificación (pronóstico y roster)", "Dashboard histórico", "Proyección anual"])

# Histórico compartido: se sube UNA vez y persiste al cambiar de vista
_hf = st.sidebar.file_uploader("Histórico único (historico.csv)", type=["csv"], key="hist")
if _hf is not None:
    st.session_state["hist_bytes"] = _hf.getvalue()
HIST = st.session_state.get("hist_bytes")

if vista == "Dashboard histórico":
    st.header("📊 Dashboard histórico por cola")
    if HIST is None:
        st.info("Sube el histórico único en la barra lateral para ver el dashboard.")
        st.stop()
    d = pd.read_csv(io.BytesIO(HIST))
    d.columns = [c.strip().lower() for c in d.columns]
    d["fecha"] = pd.to_datetime(d["fecha"], errors="coerce")
    d = d.dropna(subset=["fecha"])
    for c in ["entrantes", "atendidas", "abandonadas"]:
        d[c] = pd.to_numeric(d[c], errors="coerce").fillna(0)
    d["año"] = d["fecha"].dt.year
    d["mes"] = d["fecha"].dt.month
    d["semana"] = d["fecha"].dt.isocalendar().week.astype(int)
    d["dow"] = d["fecha"].dt.dayofweek

    f1, f2, f3, f4 = st.columns(4)
    año = f1.selectbox("Año", ["Todos"] + sorted(d["año"].unique()))
    mes = f2.selectbox("Mes", ["Todos"] + list(range(1, 13)))
    sem = f3.selectbox("Semana ISO", ["Todas"] + sorted(d["semana"].unique()))
    dias_sel = f4.multiselect("Día de semana", NOM, default=NOM)
    colas = st.multiselect("Colas", sorted(d["cola"].unique()), default=sorted(d["cola"].unique()))

    f = d.copy()
    if año != "Todos":
        f = f[f["año"] == año]
    if mes != "Todos":
        f = f[f["mes"] == mes]
    if sem != "Todas":
        f = f[f["semana"] == sem]
    f = f[f["dow"].isin([NOM.index(x) for x in dias_sel])]
    f = f[f["cola"].isin(colas)]
    if f.empty:
        st.warning("No hay datos para esos filtros.")
        st.stop()

    ent, at, ab = f["entrantes"].sum(), f["atendidas"].sum(), f["abandonadas"].sum()
    k1, k2, k3, k4 = st.columns(4)
    k1.metric("Entrantes", f"{int(ent):,}")
    k2.metric("Atendidas", f"{int(at):,}")
    k3.metric("Abandonadas", f"{int(ab):,}")
    k4.metric("% Atención", f"{at / ent * 100:.1f}%" if ent else "—")

    st.subheader("Por cola")
    by = f.groupby("cola")[["entrantes", "atendidas", "abandonadas"]].sum().reset_index()
    by["% atención"] = (by["atendidas"] / by["entrantes"].replace(0, np.nan) * 100).round(1)
    by = by.sort_values("entrantes", ascending=False)
    st.dataframe(by, use_container_width=True)

    x = np.arange(len(by)); w = 0.4
    fig, ax = plt.subplots(figsize=(10, 4))
    ax.bar(x - w / 2, by["atendidas"], w, label="Atendidas", color="#2A9D8F")
    ax.bar(x + w / 2, by["abandonadas"], w, label="Abandonadas", color="#E76F51")
    ax.set_xticks(x); ax.set_xticklabels(by["cola"], rotation=30, ha="right"); ax.legend()
    ax.set_ylabel("Llamadas")
    st.pyplot(fig)

    st.subheader("% Atención en el tiempo")
    serie = f.groupby("fecha")[["entrantes", "atendidas"]].sum()
    serie["pat"] = serie["atendidas"] / serie["entrantes"].replace(0, np.nan) * 100
    fig2, ax2 = plt.subplots(figsize=(10, 3))
    ax2.plot(serie.index, serie["pat"], color="#1F6F66", linewidth=1.5)
    ax2.axhline(96, color="#E76F51", linestyle="--", linewidth=1, label="Objetivo 96%")
    ax2.set_ylabel("% Atención"); ax2.legend()
    st.pyplot(fig2)

    if "hora" in f.columns:
        st.subheader("Por franja horaria")
        fr = f.groupby("hora")[["entrantes", "atendidas", "abandonadas"]].sum()
        fr["pat"] = fr["atendidas"] / fr["entrantes"].replace(0, np.nan) * 100
        fig3, ax3 = plt.subplots(figsize=(10, 3.5))
        ax3.bar(fr.index, fr["entrantes"], color="#C9D6D3", label="Entrantes")
        ax3.set_xlabel("Hora"); ax3.set_ylabel("Entrantes"); ax3.set_xticks(range(0, 24, 2))
        ax3b = ax3.twinx()
        ax3b.plot(fr.index, fr["pat"], color="#1F6F66", linewidth=2, label="% Atención")
        ax3b.set_ylabel("% Atención"); ax3b.set_ylim(0, 100)
        ax3.legend(loc="upper left"); ax3b.legend(loc="upper right")
        st.pyplot(fig3)
        st.caption("Barras = entrantes por hora · línea = % atención por hora.")
    st.stop()


if vista == "Proyección anual":
    st.header("📅 Proyección anual (escenarios)")
    if HIST is None:
        st.info("Sube el histórico único en la barra lateral para ver la proyección.")
        st.stop()
    d = pd.read_csv(io.BytesIO(HIST))
    d.columns = [c.strip().lower() for c in d.columns]
    d["fecha"] = pd.to_datetime(d["fecha"], errors="coerce")
    vc = "entrantes" if "entrantes" in d.columns else "vol"
    d[vc] = pd.to_numeric(d[vc], errors="coerce").fillna(0)
    d = d.dropna(subset=["fecha"])
    meses_nom = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]

    # Filtro por colas (con su % de entrantes)
    if "cola" in d.columns:
        share = d.groupby("cola")[vc].sum().sort_values(ascending=False)
        spct = (share / share.sum() * 100).round(1)
        ops = [f"{c}  ({spct[c]}%)" for c in share.index]
        sel = st.multiselect("Colas a incluir (con su % de entrantes)", ops, default=ops)
        colas_sel = [o.rsplit("  (", 1)[0] for o in sel]
        if colas_sel:
            d = d[d["cola"].isin(colas_sel)]
            st.caption(f"Incluyes el {round(share[colas_sel].sum() / share.sum() * 100, 1)}% del tráfico total.")

    dia = d.groupby(d["fecha"].dt.normalize())[vc].sum()
    if dia.empty:
        st.warning("No hay datos con esos filtros.")
        st.stop()
    años = sorted({t.year for t in dia.index})

    st.subheader("Parámetros de proyección")
    c1, c2, c3 = st.columns(3)
    año = c1.selectbox("Año a proyectar", años, index=len(años) - 1)
    scope_a = c2.selectbox("Festivos", ["Nacional España", "Cataluña (Barcelona)"])
    K = int(c3.number_input("Semanas de historia (K)", 2, 12, 4, 1))
    recencia = st.slider("Peso a lo reciente", 0.0, 0.9, 0.0, 0.1,
                         help="0 = todas las semanas pesan igual · más alto = las semanas recientes mandan")
    st.markdown("**Ajuste por mes (%)** — sube o baja cada mes proyectado de forma independiente:")
    with st.expander("Editar ajuste por mes"):
        mc = st.columns(6)
        ajuste = [mc[i % 6].number_input(meses_nom[i], -50, 100, 0, 5, key=f"am{i}") / 100.0 for i in range(12)]
    # Se guarda para que Planificación use el mismo ajuste del mes que dimensione
    st.session_state["ajuste_mes"] = {m + 1: ajuste[m] for m in range(12)}

    if scope_a == "Cataluña (Barcelona)":
        ES = holidays.Spain(years=range(min(años) - 1, max(años) + 2), subdiv="CT")
    else:
        ES = holidays.Spain(years=range(min(años) - 1, max(años) + 2))

    def fest(t):
        return t.date() in ES

    def proj_dia(t):
        wd = 6 if fest(t) else t.dayofweek
        s = dia[(dia.index < t) & (dia.index.dayofweek == wd)]
        if wd != 6:
            s = s[[not fest(x) for x in s.index]]
        s = s.tail(K)
        if len(s) == 0:
            return 0.0
        v = s.values[::-1]
        w = np.array([(1 - recencia) ** j for j in range(len(v))])
        return float((v * w).sum() / w.sum())

    # Promedios históricos
    st.subheader("Promedios históricos (colas seleccionadas)")
    g1, g2 = st.columns(2)
    pdow = dia.groupby(dia.index.dayofweek).mean().reindex(range(7)).fillna(0)
    figd, axd = plt.subplots(figsize=(5, 3))
    axd.bar([NOM[i] for i in range(7)], pdow.values, color="#2A9D8F")
    axd.set_ylabel("Entrantes/día"); axd.set_title("Promedio por día de semana")
    g1.pyplot(figd)
    dia_año = dia[dia.index.year == año]
    sem = dia_año.resample("W").sum()
    import matplotlib.dates as mdates
    figs, axs = plt.subplots(figsize=(5, 3))
    axs.plot(sem.index, sem.values, color="#1F6F66", marker="o", markersize=3)
    axs.set_ylabel("Entrantes/semana"); axs.set_title(f"Volumen por semana {año}")
    axs.grid(True, alpha=0.3)
    axs.xaxis.set_major_locator(mdates.MonthLocator())
    axs.xaxis.set_major_formatter(mdates.DateFormatter("%b"))
    g2.pyplot(figs)

    # Proyección mensual (real intacto; multiplicador y % solo a lo proyectado)
    filas = []
    for mth in range(1, 13):
        ini = pd.Timestamp(year=año, month=mth, day=1)
        fin = ini + pd.offsets.MonthEnd(0)
        dias_mes = pd.date_range(ini, fin)
        real_part = proj_part = 0.0
        for t in dias_mes:
            if t in dia.index:
                real_part += float(dia.loc[t])
            else:
                proj_part += proj_dia(t)
        n_real = sum(1 for t in dias_mes if t in dia.index)
        estado = "REAL" if n_real >= len(dias_mes) else ("EN CURSO" if n_real > 0 else "PROYECTADO")
        total = real_part + proj_part * (1 + ajuste[mth - 1])
        filas.append({"Mes": meses_nom[mth - 1], "Estado": estado,
                      "Real a la fecha": int(round(real_part)), "Proyectado mes": int(round(total))})
    tab = pd.DataFrame(filas)
    m1, m2 = st.columns(2)
    m1.metric(f"Proyección total {año}", f"{int(tab['Proyectado mes'].sum()):,}")
    m2.metric("Real acumulado", f"{int(tab['Real a la fecha'].sum()):,}")
    st.dataframe(tab, use_container_width=True)
    cores = {"REAL": "#2A9D8F", "EN CURSO": "#E9C46A", "PROYECTADO": "#C9D6D3"}
    figp, axp = plt.subplots(figsize=(11, 4))
    axp.bar(tab["Mes"], tab["Proyectado mes"], color=[cores[e] for e in tab["Estado"]])
    axp.set_ylabel("Llamadas"); axp.set_title(f"Volumen mensual {año} (real + proyectado)")
    st.pyplot(figp)
    st.caption("Verde = real · Amarillo = en curso · Gris = proyectado. El ajuste por mes se aplica solo a los días proyectados y se usa también en Planificación.")
    st.stop()


# ---------------- Parámetros ----------------
st.sidebar.header("Parámetros")
AHT = st.sidebar.number_input("AHT (seg)", 60, 1200, 420, 10)
SLA = st.sidebar.slider("SLA objetivo", 0.50, 0.99, 0.80, 0.01)
ASA = st.sidebar.number_input("ASA objetivo (seg)", 5, 120, 20, 5)
OCC = st.sidebar.slider("Ocupación tope (OCC)", 0.40, 0.95, 0.65, 0.01)
UTL = st.sidebar.slider("Utilización (UTL)", 0.60, 1.00, 0.88, 0.01)
ABS = st.sidebar.slider("Absentismo", 0.00, 0.40, 0.15, 0.01)
ESP_MAX = st.sidebar.number_input("Agentes España (máx, solo L–V)", 0, 200, 12, 1)
LARGO = int(st.sidebar.number_input("Duración turno (h)", 6, 12, 9, 1))
NDA_OBJ = st.sidebar.slider("NDA objetivo (nivel de atención)", 0.80, 0.999, 0.96, 0.005)
PACIENCIA = st.sidebar.number_input("Paciencia media (seg)", 20, 600, 90, 10)


# ---------------- Núcleo: dimensionar + roster ----------------
def dimension_roster(largo, AHT, SLA, ASA, OCC, UTL, ESP_MAX, LARGO, NDA_OBJ, PACIENCIA):
    largo = largo.copy()
    largo["dow"] = pd.to_datetime(largo["fecha"]).dt.dayofweek
    volmax = {dw: [0.0] * 24 for dw in range(7)}
    for (dw, h), g in largo.groupby(["dow", "intervalo"]):
        volmax[dw][int(h)] = float(g["volumen"].max())
    peak = {dw: [0] * 24 for dw in range(7)}
    occ = {dw: [0.0] * 24 for dw in range(7)}
    nda = {dw: [1.0] * 24 for dw in range(7)}
    for dw in range(7):
        for h in range(24):
            ca = volmax[dw][h] * AHT / 3600
            if ca <= 0:
                continue
            a = int(ca) + 1
            while nivel_servicio(a, ca, AHT, ASA) < SLA:
                a += 1
            en = max(a, math.ceil(ca / OCC))
            while nivel_atencion(en, ca, AHT, PACIENCIA) < NDA_OBJ:   # piso por NDA
                en += 1
            peak[dw][h] = math.ceil(en / UTL)
            occ[dw][h] = ca / en
            nda[dw][h] = nivel_atencion(en, ca, AHT, PACIENCIA)

    H = 24
    turnos = {ini: [(ini + k) % H for k in range(LARGO)] for ini in range(H)}
    m = cp_model.CpModel()
    xc = {(t, p): m.NewIntVar(0, 300, f"c{t}_{p}") for t in turnos for p in range(7)}
    xe = {t: m.NewIntVar(0, 300, f"e{t}") for t in turnos}
    for dw in range(7):
        for h in range(H):
            col = sum(xc[(t, p)] for t in turnos for p in range(7) if dw not in LIBRES[p] and h in turnos[t])
            esp = sum(xe[t] for t in turnos if dw not in LIBRES[5] and h in turnos[t])
            m.Add(col + esp >= peak[dw][h])
    m.Add(sum(xe.values()) <= int(ESP_MAX))
    m.Minimize(sum(xc.values()) * 100 - sum(xe.values()))
    sv = cp_model.CpSolver(); sv.Solve(m)

    xe_s = {str(t): sv.Value(xe[t]) for t in turnos if sv.Value(xe[t]) > 0}
    xc_s = {f"{t}_{p}": sv.Value(xc[(t, p)]) for t in turnos for p in range(7) if sv.Value(xc[(t, p)]) > 0}
    te = sum(xe_s.values()); tc = sum(xc_s.values())

    def crew(req):
        mm = cp_model.CpModel()
        xx = {t: mm.NewIntVar(0, 200, str(t)) for t in turnos}
        for h in range(H):
            mm.Add(sum(xx[t] for t in turnos if h in turnos[t]) >= req[h])
        mm.Minimize(sum(xx.values()))
        ss = cp_model.CpSolver(); ss.Solve(mm)
        return sum(ss.Value(v) for v in xx.values())
    cw = max(crew(peak[5]), crew(peak[6]))
    vals = [nda[dw][h] for dw in range(7) for h in range(24) if peak[dw][h] > 0]
    return {"peak": peak, "occ": occ, "nda": nda, "nda_min": min(vals) if vals else 1.0,
            "xe": xe_s, "xc": xc_s, "te": te, "tc": tc,
            "crew": cw, "rot": cw * 4 // 2, "total": int(largo["volumen"].sum())}


# ---------------- Producir 'largo' (dos orígenes) ----------------
def largo_desde_crosstab(file_bytes):
    raw = pd.read_excel(io.BytesIO(file_bytes))
    fechas = raw.iloc[0]; first = raw.columns[0]
    d = raw.drop(index=0)
    d = d[d[first].astype(str).str.lower() != "total"].rename(columns={first: "intervalo"})
    L = d.melt(id_vars="intervalo", var_name="col", value_name="volumen")
    L["fecha"] = pd.to_datetime(L["col"].map(fechas), errors="coerce")
    L = L.dropna(subset=["fecha"])
    L["volumen"] = pd.to_numeric(L["volumen"], errors="coerce")
    L = L.dropna(subset=["volumen"])

    def hh(x):
        try:
            return int(x)
        except Exception:
            return int(str(x).split(":")[0])
    L["intervalo"] = L["intervalo"].apply(hh)
    L = L[(L["intervalo"] >= 0) & (L["intervalo"] <= 23)]
    return L[["fecha", "intervalo", "volumen"]]


def largo_desde_historico(file_bytes, mes, scope, K, semanas, ajuste=0.0):
    # Archivo ÚNICO (CSV: fecha, hora, cola, entrantes, atendidas, abandonadas)
    df = pd.read_csv(io.BytesIO(file_bytes))
    df.columns = [c.strip().lower() for c in df.columns]
    df["fecha"] = pd.to_datetime(df["fecha"], errors="coerce")
    df["hora"] = pd.to_numeric(df["hora"], errors="coerce")
    volcol = "entrantes" if "entrantes" in df.columns else "vol"
    df[volcol] = pd.to_numeric(df[volcol], errors="coerce").fillna(0)
    df = df.dropna(subset=["fecha", "hora"])
    df["hora"] = df["hora"].astype(int)
    hh = df.groupby(["fecha", "hora"])[volcol].sum().reset_index()
    hh.columns = ["fecha", "hora", "vol"]
    dia = hh.groupby("fecha")["vol"].sum()

    if scope == "Cataluña (Barcelona)":
        ES = holidays.Spain(years=range(2023, 2028), subdiv="CT")
    else:
        ES = holidays.Spain(years=range(2023, 2028))

    def fest(t):
        return t.date() in ES

    def total_diario(t):
        wd = 6 if fest(t) else t.dayofweek
        s = dia[(dia.index < t) & (dia.index.dayofweek == wd)]
        if wd != 6:
            s = s[[not fest(d) for d in s.index]]
        return s.tail(K).mean()

    ini = pd.Timestamp(mes + "-01")
    win = hh[(hh["fecha"] < ini) & (hh["fecha"] >= ini - pd.Timedelta(weeks=semanas))]
    win = win[[not fest(d) for d in win["fecha"]]]
    perfil = win.groupby([win["fecha"].dt.dayofweek, "hora"])["vol"].mean().unstack(fill_value=0)
    perfil = perfil.div(perfil.sum(axis=1), axis=0)

    filas = []
    for t in pd.date_range(ini, ini + pd.offsets.MonthEnd(0)):
        dt = total_diario(t)
        wd = 6 if fest(t) else t.dayofweek
        for hr in range(24):
            filas.append({"fecha": t, "intervalo": hr, "volumen": round(dt * perfil.loc[wd].get(hr, 0) * (1 + ajuste))})
    return pd.DataFrame(filas)


@st.cache_data(show_spinner="Dimensionando y optimizando…")
def resolver_crosstab(file_bytes, AHT, SLA, ASA, OCC, UTL, ESP_MAX, LARGO, NDA_OBJ, PACIENCIA):
    return dimension_roster(largo_desde_crosstab(file_bytes), AHT, SLA, ASA, OCC, UTL, ESP_MAX, LARGO, NDA_OBJ, PACIENCIA)


@st.cache_data(show_spinner="Pronosticando, dimensionando y optimizando…")
def resolver_historico(file_bytes, mes, scope, K, semanas, AHT, SLA, ASA, OCC, UTL, ESP_MAX, LARGO, NDA_OBJ, PACIENCIA, ajuste=0.0):
    return dimension_roster(largo_desde_historico(file_bytes, mes, scope, K, semanas, ajuste),
                            AHT, SLA, ASA, OCC, UTL, ESP_MAX, LARGO, NDA_OBJ, PACIENCIA)


def turnos_dict(largo):
    return {ini: [(ini + k) % 24 for k in range(largo)] for ini in range(24)}


def cubierto(S, turnos, dw, h):
    cob = sum(v for k, v in S["xc"].items()
              if dw not in LIBRES[int(k.split("_")[1])] and h in turnos[int(k.split("_")[0])])
    cob += sum(v for k, v in S["xe"].items()
               if dw not in LIBRES[5] and h in turnos[int(k)])
    return cob


def build_excel(S, turnos, params):
    AHT, SLA, ASA, OCC, UTL, ABS = params
    FT = "Arial"
    HEAD = PatternFill("solid", start_color="1F6F66")
    HF = Font(name=FT, bold=True, color="FFFFFF")
    TITLE = Font(name=FT, bold=True, size=14, color="1F6F66")
    BOLD = Font(name=FT, bold=True); REG = Font(name=FT)
    CEN = Alignment(horizontal="center")
    thin = Side(style="thin", color="D9D9D9"); BORD = Border(thin, thin, thin, thin)

    def hrow(ws, row, cols):
        for j, c in enumerate(cols, 1):
            x = ws.cell(row, j, c); x.fill = HEAD; x.font = HF; x.alignment = CEN; x.border = BORD

    wb = Workbook()
    te, tc = S["te"], S["tc"]; peak = S["peak"]

    ws = wb.active; ws.title = "Resumen"
    ws["A1"] = "Roster de dimensionamiento"; ws["A1"].font = TITLE
    r = 3; ws.cell(r, 1, "PARÁMETROS").font = BOLD; r += 1
    for k, v, f in [("AHT (seg)", AHT, "0"), ("SLA", SLA, "0%"), ("ASA (seg)", ASA, "0"),
                    ("OCC", OCC, "0%"), ("UTL", UTL, "0%"), ("Absentismo", ABS, "0%")]:
        ws.cell(r, 1, k).font = REG; c = ws.cell(r, 2, v); c.font = REG; c.number_format = f; r += 1
    r += 1; ws.cell(r, 1, "PLANTILLA").font = BOLD; r += 1
    for k, v in [("España (L–V)", te), ("Colombia", tc), ("TOTAL presentes", te + tc),
                 ("En nómina (+absentismo)", math.ceil((te + tc) / (1 - ABS)))]:
        b = "TOTAL" in k or "nómina" in k
        ws.cell(r, 1, k).font = BOLD if b else REG
        ws.cell(r, 2, v).font = BOLD if b else REG; r += 1
    r += 1; hrow(ws, r, ["Día", "Trabajan", "Pico requerido"]); r += 1
    for dw in range(7):
        trab = sum(v for k, v in S["xc"].items() if dw not in LIBRES[int(k.split("_")[1])])
        trab += sum(v for k, v in S["xe"].items() if dw not in LIBRES[5])
        for j, v in enumerate([NOM[dw], trab, max(peak[dw])], 1):
            x = ws.cell(r, j, v); x.font = REG; x.border = BORD
            if j > 1:
                x.alignment = CEN
        r += 1
    ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 14; ws.column_dimensions["C"].width = 16

    ws = wb.create_sheet("Plan_Turnos")
    hrow(ws, 1, ["País", "Inicio", "Fin", "Cantidad", "Días libres"])
    for w, col in zip([12, 9, 9, 11, 18], "ABCDE"):
        ws.column_dimensions[col].width = w
    r = 2
    L = len(turnos[0])

    def add(r, pais, t, cant, off):
        o = sorted(off)
        vals = [pais, f"{t:02d}:00", f"{(t + L) % 24:02d}:00", cant, f"{NOM[o[0]]}, {NOM[o[1]]}"]
        for j, v in enumerate(vals, 1):
            x = ws.cell(r, j, v); x.font = REG; x.border = BORD
            if j in (2, 3, 4):
                x.alignment = CEN
            if pais == "España":
                x.fill = PatternFill("solid", start_color="FCE4D6")
        return r + 1
    for k in sorted(S["xe"], key=int):
        r = add(r, "España", int(k), S["xe"][k], {5, 6})
    for k in sorted(S["xc"], key=lambda x: (int(x.split("_")[0]), int(x.split("_")[1]))):
        t, p = map(int, k.split("_")); r = add(r, "Colombia", t, S["xc"][k], LIBRES[p])
    ws.cell(r + 1, 1, "TOTAL").font = BOLD; ws.cell(r + 1, 4, f"=SUM(D2:D{r-1})").font = BOLD

    wsg = wb.create_sheet("Grafico")
    wsg["A1"] = "Programados vs Requerido — por día"; wsg["A1"].font = TITLE
    hdr = ["Hora"]
    for dw in range(7):
        hdr += [f"{NOM[dw]} Req", f"{NOM[dw]} Cub"]
    for j, c in enumerate(hdr, 1):
        x = wsg.cell(3, j, c); x.fill = HEAD; x.font = HF; x.alignment = CEN
    for h in range(24):
        wsg.cell(4 + h, 1, f"{h:02d}:00").font = REG
        for dw in range(7):
            wsg.cell(4 + h, 2 + dw * 2, peak[dw][h]).font = REG
            wsg.cell(4 + h, 3 + dw * 2, cubierto(S, turnos, dw, h)).font = REG
    wsg.column_dimensions["A"].width = 8
    arow = 30
    for dw in range(7):
        bar = BarChart(); bar.type = "col"; bar.title = f"{NOM[dw]} — programados vs requerido"
        bar.height = 6.5; bar.width = 20; bar.y_axis.title = "Agentes"
        bar.add_data(Reference(wsg, min_col=2 + dw * 2, min_row=3, max_row=27), titles_from_data=True)
        bar.set_categories(Reference(wsg, min_col=1, min_row=4, max_row=27))
        bar.series[0].graphicalProperties.solidFill = "C9D6D3"
        ln = LineChart()
        ln.add_data(Reference(wsg, min_col=3 + dw * 2, min_row=3, max_row=27), titles_from_data=True)
        ln.series[0].graphicalProperties.line.solidFill = "1F6F66"
        ln.series[0].graphicalProperties.line.width = 22000
        ln.series[0].smooth = False
        bar += ln
        wsg.add_chart(bar, f"A{arow}"); arow += 14
    wsg.sheet_view.showGridLines = False

    bio = io.BytesIO(); wb.save(bio); return bio.getvalue()


# ================== UI ==================
st.subheader("1) Origen de la demanda")
modo = st.radio("¿Cómo obtenemos el pronóstico?",
                ["Generar desde histórico", "Subir pronóstico ya hecho"])

S = None
if modo == "Generar desde histórico":
    st.caption("Usa el mismo archivo del Dashboard y la Proyección anual (se sube una vez en la barra lateral). "
               "Se genera en Colab a partir del histórico crudo.")
    c1, c2, c3, c4 = st.columns(4)
    mes = c1.text_input("Mes (AAAA-MM)", "2026-06")
    scope = c2.selectbox("Festivos", ["Nacional España", "Cataluña (Barcelona)"])
    K = int(c3.number_input("Semanas promedio", 2, 12, 4, 1))
    semanas = int(c4.number_input("Ventana perfil (sem.)", 2, 12, 6, 1))
    # Ajuste del mes definido en la Proyección anual (si existe)
    try:
        mnum = int(mes.split("-")[1])
    except Exception:
        mnum = 0
    ajuste_mes = st.session_state.get("ajuste_mes", {}).get(mnum, 0.0)
    if ajuste_mes:
        st.caption(f"Aplicando ajuste de {int(ajuste_mes * 100):+d}% definido para ese mes en la Proyección anual.")
    if HIST is None:
        st.info("Sube el histórico único en la barra lateral.")
    else:
        try:
            S = resolver_historico(HIST, mes, scope, K, semanas, AHT, SLA, ASA, OCC, UTL, ESP_MAX, LARGO, NDA_OBJ, PACIENCIA, ajuste_mes)
        except Exception as e:
            st.error(f"No pude procesar el histórico: {e}")
else:
    up = st.file_uploader("Sube tu pronóstico (días en columnas, intervalos en filas)", type=["xlsx", "xls"])
    if up is not None:
        S = resolver_crosstab(up.getvalue(), AHT, SLA, ASA, OCC, UTL, ESP_MAX, LARGO, NDA_OBJ, PACIENCIA)

if S is None:
    st.info("Sube el archivo para continuar.")
    st.stop()

turnos = turnos_dict(LARGO)

st.subheader("2) Plantilla requerida")
c1, c2, c3, c4 = st.columns(4)
c1.metric("España (L–V)", S["te"])
c2.metric("Colombia", S["tc"])
c3.metric("TOTAL presentes", S["te"] + S["tc"])
c4.metric("En nómina (+absentismo)", math.ceil((S["te"] + S["tc"]) / (1 - ABS)))
st.caption(f"Volumen del mes: {S['total']:,} llamadas · Rotación findes Colombia: cuadrilla {S['crew']}/finde, "
           f"{S['rot']} en rotación (2 grupos de {S['crew']}, máx. 2 findes/mes).")
st.caption(f"NDA previsto (nivel de atención): mínimo {S['nda_min']:.1%} · objetivo {NDA_OBJ:.0%}. "
           f"El tope de OCC ya garantiza un NDA alto; solo aprieta en horas de muy bajo volumen.")

st.subheader("3) Programado vs Requerido y Ocupación")
dsel = st.selectbox("Día", list(range(7)), format_func=lambda i: NOM[i])
req = [S["peak"][dsel][h] for h in range(24)]
cob = [cubierto(S, turnos, dsel, h) for h in range(24)]
fig, ax = plt.subplots(figsize=(10, 3.5))
ax.bar(range(24), req, color="#C9D6D3", label="Requerido")
ax.step(range(24), cob, where="mid", color="#1F6F66", linewidth=2, label="Programado")
ax.set_xlabel("Hora"); ax.set_ylabel("Agentes"); ax.set_xticks(range(0, 24, 2)); ax.legend()
st.pyplot(fig)

st.markdown("**Ocupación por hora**")
occv = [S["occ"][dsel][h] * 100 for h in range(24)]
fig2, ax2 = plt.subplots(figsize=(10, 3))
colors = ["#E76F51" if o < 30 else ("#E9C46A" if o < OCC * 100 * 0.85 else "#2A9D8F") for o in occv]
ax2.bar(range(24), occv, color=colors)
ax2.axhline(OCC * 100, color="#264653", linestyle="--", linewidth=1, label=f"Objetivo OCC {OCC:.0%}")
ax2.set_xlabel("Hora"); ax2.set_ylabel("Ocupación %"); ax2.set_xticks(range(0, 24, 2))
ax2.set_ylim(0, 100); ax2.legend()
st.pyplot(fig2)
st.caption("Rojo = horas de baja ocupación (poca demanda). La línea marca tu objetivo de OCC.")

st.markdown("**Nivel de atención (NDA) por hora**")
ndav = [S["nda"][dsel][h] * 100 for h in range(24)]
cols3 = ["#E76F51" if (S["nda"][dsel][h] < NDA_OBJ and S["peak"][dsel][h] > 0) else "#2A9D8F" for h in range(24)]
fig3, ax3 = plt.subplots(figsize=(10, 3))
ax3.bar(range(24), ndav, color=cols3)
ax3.axhline(NDA_OBJ * 100, color="#264653", linestyle="--", linewidth=1, label=f"Objetivo {NDA_OBJ:.0%}")
ax3.set_xlabel("Hora"); ax3.set_ylabel("NDA %"); ax3.set_xticks(range(0, 24, 2))
ax3.set_ylim(0, 100); ax3.legend()
st.pyplot(fig3)
st.caption("Nivel de atención previsto (Erlang con abandono). Rojo = por debajo del objetivo.")

st.subheader("4) Plan de turnos")
filas = []
for k in sorted(S["xe"], key=int):
    t = int(k); filas.append(["España", f"{t:02d}:00", f"{(t+LARGO)%24:02d}:00", S["xe"][k], "Sáb, Dom"])
for k in sorted(S["xc"], key=lambda x: (int(x.split("_")[0]), int(x.split("_")[1]))):
    t, p = map(int, k.split("_")); o = sorted(LIBRES[p])
    filas.append(["Colombia", f"{t:02d}:00", f"{(t+LARGO)%24:02d}:00", S["xc"][k], f"{NOM[o[0]]}, {NOM[o[1]]}"])
st.dataframe(pd.DataFrame(filas, columns=["País", "Inicio", "Fin", "Cantidad", "Días libres"]),
             use_container_width=True)

st.subheader("5) Comparar con tu plantilla actual (opcional)")
ag = st.file_uploader("Sube tu Excel de agentes (columnas MODO, CENTRO, ESTADO…)",
                      type=["xlsx", "xls"], key="agentes")
if ag is not None:
    A = pd.read_excel(ag)
    if "ESTADO" in A.columns:
        A = A[A["ESTADO"].astype(str).str.upper() == "ACTIVO"]
    if "MODO" in A.columns:
        modos = sorted(A["MODO"].dropna().astype(str).unique())
        default = ["MULTISKILL"] if "MULTISKILL" in modos else modos
        sel = st.multiselect("Skill(s) a contar (MODO)", modos, default=default)
        if sel:
            A = A[A["MODO"].astype(str).isin(sel)]
    if "CENTRO" in A.columns:
        A["pais"] = A["CENTRO"].astype(str).str.upper().apply(
            lambda c: "España" if c in ("SEVILLA", "BARCELONA") else "Colombia")
        act_e = int((A["pais"] == "España").sum())
        act_c = int((A["pais"] == "Colombia").sum())
    else:
        act_e, act_c = 0, len(A)

    def nomina(x):
        return math.ceil(x / (1 - ABS))

    comp = pd.DataFrame({
        "País": ["España", "Colombia", "TOTAL"],
        "Requerido (presentes)": [S["te"], S["tc"], S["te"] + S["tc"]],
        "Necesarios en nómina (+abs)": [nomina(S["te"]), nomina(S["tc"]), nomina(S["te"]) + nomina(S["tc"])],
        "Agentes actuales": [act_e, act_c, act_e + act_c],
    })
    comp["Gap (actual − nómina)"] = comp["Agentes actuales"] - comp["Necesarios en nómina (+abs)"]
    st.dataframe(comp, use_container_width=True)
    g = int(comp.iloc[2]["Gap (actual − nómina)"])
    if g >= 0:
        st.success(f"Tienes {g} agentes de margen sobre lo necesario en nómina.")
    else:
        st.warning(f"Te faltan {-g} agentes respecto a lo necesario en nómina.")
    st.caption("Centros: SEVILLA y BARCELONA → España; el resto → Colombia.")

st.subheader("6) Descargar")
st.download_button("⬇️ Roster en Excel (con gráficos)",
                   build_excel(S, turnos, (AHT, SLA, ASA, OCC, UTL, ABS)),
                   file_name="roster.xlsx",
                   mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
